from pathlib import Path
import json
import fitz
import pandas as pd
import geopandas as gpd
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'platform' / 'work' / 'sources'
OUT.mkdir(parents=True, exist_ok=True)
for path in sorted((ROOT / 'data').rglob('*.pdf')):
    doc = fitz.open(path)
    pages = [{'pdf_page': i + 1, 'text': p.get_text()} for i, p in enumerate(doc)]
    (OUT / (path.stem + '.json')).write_text(json.dumps(pages, ensure_ascii=False), encoding='utf-8')
    print(path.name, len(pages), 'pages', sum(len(p['text']) for p in pages), 'characters', flush=True)
for path in sorted((ROOT / 'data').rglob('*.docx')):
    doc = Document(path)
    text = '\n'.join(p.text for p in doc.paragraphs)
    text += '\n' + '\n'.join(' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows)
    (OUT / (path.stem + '.txt')).write_text(text, encoding='utf-8')
    print(path.name, text, flush=True)
plots = gpd.read_file(ROOT / 'data/fira/plots/Parcelas_Reto_AGC_CONJUNTO.shp')
print('PLOTS', plots.crs, plots.drop(columns='geometry').head().to_string(), plots.columns.tolist(), flush=True)
for path in sorted((ROOT / 'data/fira/CSV-files').glob('*.csv')):
    df = pd.read_csv(path)
    print(path.name, df.shape, 'duplicates', df.duplicated().sum(), flush=True)
    if 'fecha_captura' in df:
        dates = pd.to_datetime(df.fecha_captura, dayfirst=True)
        print('DATES', dates.min(), dates.max(), 'SENSORS', df.sensor.value_counts().to_dict(), flush=True)
        print('CLOUDS', df.groupby('sensor').porcentaje_nubosidad.agg(['min','max','count']).to_dict(), flush=True)
        print('CLOUD ZERO', df.loc[df.porcentaje_nubosidad.eq(0)].groupby('sensor').ndvi_promedio.agg(['count','mean']).to_dict(), flush=True)
        print('CLOUD 100', df.loc[df.porcentaje_nubosidad.eq(100)].groupby('sensor').ndvi_promedio.agg(['count','mean']).to_dict(), flush=True)
